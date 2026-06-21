import subprocess, sys, json

# Install openpyxl if needed - already installed
# Check what SheetJS sees by reading the cached formula value AND cell.w equivalent

# Read with openpyxl data_only to get cached values
import openpyxl
wb = openpyxl.load_workbook(
    r'c:\proyecto power apps\comparar excels\Portafolio_Demanda_TI_v12 (6).xlsm',
    read_only=True, keep_vba=True, data_only=True
)
ws = wb['Demanda Tactica']
headers = [cell.value for cell in list(ws.iter_rows(min_row=5, max_row=5, values_only=False))[0]]
pct_cols = {i: h for i, h in enumerate(headers) if h in ['% PLANIFICADO','% COMPLETADO','GAP','RATIO']}
print('PCT columns:')
for i, h in pct_cols.items():
    print(f'  col {i}: {h}')

print('\nFirst 5 data rows (cached values):')
for row in ws.iter_rows(min_row=6, max_row=10, values_only=False):
    vals = {}
    for i, h in pct_cols.items():
        cell = row[i]
        vals[h] = {
            'value': cell.value,
            'number_format': cell.number_format,
        }
    print(vals)

# Now check what number_format those cells have
print('\nNumber formats:')
for row in ws.iter_rows(min_row=6, max_row=7, values_only=False):
    for i, h in pct_cols.items():
        cell = row[i]
        print(f'  {h} row{cell.row}: value={cell.value!r}  format={cell.number_format!r}')
